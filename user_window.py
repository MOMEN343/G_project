import os
import sys
import shutil
from docx import Document
from datetime import datetime, date, time, timedelta
from db import DataBase
from PyQt5 import uic, QtWidgets, QtCore
from PyQt5.QtWidgets import (
    QMainWindow, QWidget, QVBoxLayout, QLabel, QPushButton, QScrollArea, 
    QGridLayout, QHBoxLayout, QCheckBox, QMessageBox, QMenu, 
    QWidgetAction, QFrame, QHeaderView
)
from PyQt5.QtGui import QColor, QFontDatabase
from PyQt5.QtCore import Qt, QPoint, QTimer, QTime

def resource_path(relative_path):
    try:
        base_path = sys._MEIPASS
    except Exception:
        base_path = os.path.abspath(".")
    return os.path.join(base_path, relative_path)

class UserWindow(QMainWindow):
    def __init__(self, current_user_id, main_shell=None):
        super().__init__()
        self.current_user_id = current_user_id
        self.main_shell = main_shell
        self.db = DataBase()

        # Load the new UI
        uic.loadUi("employee.ui", self)

        self.setStyleSheet("""
        * {
            font-family: "Alyamama", "Segoe UI Symbol";
            color: #452829;
        }
        QLineEdit {
            border: 1px solid #ccc;
            border-radius: 5px;
            padding: 5px;
            background-color: white;
        }
        QLineEdit:focus {
            border: 1px solid #452829;
            background-color: #fcfcfc;
        }
        
        /* Premium QCalendarWidget Styling */
        QCalendarWidget QWidget {
            background-color: white;
        }
        QCalendarWidget #qt_calendar_navigationbar {
            background-color: #452829;
            border-top-left-radius: 10px;
            border-top-right-radius: 10px;
        }
        QCalendarWidget QToolButton {
            color: white;
            font-weight: bold;
            font-family: 'Alyamama';
            font-size: 15px;
            icon-size: 20px;
            padding: 5px;
            background: transparent;
            border: none;
        }
        QCalendarWidget QToolButton::menu-indicator {
            image: none;
        }
        QCalendarWidget QToolButton:hover {
            background-color: rgba(176, 141, 87, 0.5);
            border-radius: 8px;
        }
        QCalendarWidget QToolButton#qt_calendar_prevmonth {
            qproperty-icon: none;
            qproperty-text: "◀";
        }
        QCalendarWidget QToolButton#qt_calendar_nextmonth {
            qproperty-icon: none;
            qproperty-text: "▶";
        }
        QCalendarWidget QAbstractItemView:enabled {
            color: #452829;
            font-size: 14px;
            selection-background-color: #b08d57;
            selection-color: white;
            outline: none;
            border: none;
            background-color: white;
        }
        QCalendarWidget QAbstractItemView:disabled {
            color: #bbb;
        }
        QCalendarWidget #qt_calendar_calendarview {
            border: 1px solid #e8e0d8;
            border-bottom-left-radius: 10px;
            border-bottom-right-radius: 10px;
        }
        QCalendarWidget QSpinBox {
            background-color: white;
            color: #452829;
            selection-background-color: #b08d57;
            border: none;
            font-size: 14px;
            margin: 0px;
        }
        QCalendarWidget QSpinBox::up-button, QCalendarWidget QSpinBox::down-button {
            width: 0px;
        }
        
        """)

        # Connect Buttons
        self.add_case.clicked.connect(self.show_add_case)
        self.docments.clicked.connect(self.show_documents)
        self.logoutBtn.clicked.connect(self.log_out)
        self.master_record.clicked.connect(self.show_master_record)
        self.btn_scheduling.clicked.connect(self.show_scheduling)
        self.btn_save_session.clicked.connect(self.save_session)
        self.calendar.clicked.connect(self.show_calendar)
        self.searchMasterRecord.textChanged.connect(self.filter_master_record)
        self.searchScheduling.textChanged.connect(self.filter_scheduling)
        
        if hasattr(self, 'btn_edit_petition'):
            self.btn_edit_petition.clicked.connect(self.edit_selected_petition)
        if hasattr(self, 'btn_delete_petition'):
            self.btn_delete_petition.clicked.connect(self.delete_selected_petitions)
        
        
        # 2. Fix Labels and Swap with Spacers to ensure Right Side
        for name in ['labelJudgeWrapper', 'labelDateWrapper', 'labelTimeWrapper']:
            layout = self.findChild(QtWidgets.QHBoxLayout, name)
            if layout:
                # Ensure Label is on the RIGHT, Spacer on the LEFT
                # Qt's QHBoxLayout in RTL mode puts the first item added on the right.
                layout.setDirection(QtWidgets.QBoxLayout.LeftToRight) # Set base LTR
                items = []
                while layout.count():
                    items.append(layout.takeAt(0))
                
                # Order: [Spacer, Label] -> In LTR this makes Label on the right.
                label_item = None
                spacer_item = None
                for it in items:
                    if it.widget() and isinstance(it.widget(), QtWidgets.QLabel):
                        label_item = it
                    elif it.spacerItem():
                        spacer_item = it
                
                if spacer_item: layout.addItem(spacer_item)
                if label_item: layout.addItem(label_item)
                if label_item and label_item.widget():
                    label_item.widget().setStyleSheet("color: #b08d57; font-size: 14px; font-weight: bold; font-family: 'Alyamama';")
                    label_item.widget().setAlignment(Qt.AlignRight | Qt.AlignVCenter)

        # 3. Fix main card layout direction (Items: Judge, Date, Time, SaveBtn)
        # Should be RTL: Judge(Right) -> Date -> Time -> SaveBtn(Left)
        if hasattr(self, 'bottomCardLayout'):
            self.bottomCardLayout.setDirection(QtWidgets.QBoxLayout.RightToLeft)
            self.schedulingBottomCard.setLayoutDirection(Qt.LeftToRight) # Keep the card itself LTR for proper button placement
        
        if hasattr(self, 'notification'):
            self.notification.clicked.connect(self.show_notifications)
            self.notification.setFocusPolicy(Qt.NoFocus)
        
        # --- Notification Badge ---
        if hasattr(self, 'notification') and hasattr(self, 'badge_label'):
            self.badge_label.setParent(self.notification)
            self.badge_label.move(24, 2)
            
            self.timer = QTimer(self)
            self.timer.timeout.connect(self.update_badge)
            self.timer.start(5000)
            self.update_badge()

        # --- Deletion Feature UI Setup ---
        self.selected_documents = set()
        self.doc_checkboxes = [] 
        
        if hasattr(self, 'check_all_docs'):
            self.check_all_docs.stateChanged.connect(self.select_all_documents)


        # Ensure we start at the empty page
        if hasattr(self, 'mainStack'):
             self.mainStack.setCurrentWidget(self.page_empty)
             
        if hasattr(self, 'files_grid'):
            self.files_grid.setAlignment(Qt.AlignLeft | Qt.AlignTop)

        if hasattr(self, 'page_calendar') and hasattr(self, 'verticalLayout_calendar'):
            for w in self.page_calendar.findChildren(QWidget):
                if w.objectName() in ["calendarLeftPanel", "calendarHeader", "search_calendar", "btn_calendar_add", "label_calendar_date", "footerLayout"] or \
                   isinstance(w, (QPushButton, QtWidgets.QLineEdit, QLabel)) and w.parent() == self.page_calendar:
                    w.hide()
            
            while self.verticalLayout_calendar.count():
                child = self.verticalLayout_calendar.takeAt(0)
                if child.widget(): child.widget().hide()

            self.page_calendar.setStyleSheet("QWidget#page_calendar { background-color: transparent; }")

            self.header_card = QWidget()
            self.header_card.setFixedHeight(85)
            h_layout = QHBoxLayout(self.header_card)
            h_layout.setContentsMargins(30, 20, 30, 10)
            
            title_lbl = QLabel("🗄 جدول القضاة")
            title_lbl.setStyleSheet("color: #452829; font-size: 20pt; font-weight: bold;")
            
            self.date_nav_box = QWidget()
            d_layout = QHBoxLayout(self.date_nav_box)
            d_layout.setSpacing(15)
            btn_style = "QPushButton { color: #b08d57; font-size: 22px; border: none; background: transparent; font-weight: bold; } QPushButton:hover { color: #452829; }"
            
            self.custom_btn_prev = QPushButton("◀")
            self.custom_btn_prev.setStyleSheet(btn_style)
            self.custom_label_date = QLabel()
            self.custom_label_date.setStyleSheet("color: #452829; font-size: 19px; font-family: 'Alyamama';")
            self.custom_btn_next = QPushButton("▶")
            self.custom_btn_next.setStyleSheet(btn_style)
            
            d_layout.addWidget(self.custom_btn_next)
            d_layout.addWidget(self.custom_label_date)
            d_layout.addWidget(self.custom_btn_prev)

            self.custom_search = QtWidgets.QLineEdit()
            self.custom_search.setPlaceholderText("ابحث في الجدول... 🔍")
            self.custom_search.setFixedWidth(350)
            self.custom_search.setFixedHeight(45)
            self.custom_search.setStyleSheet("""
                QLineEdit {
                    background-color: white;
                    border: 2px solid #e0e0e0;
                    border-radius: 10px;
                    padding: 5px 15px;
                    font-size: 14px;
                    color: #452829;
                }
                QLineEdit:focus {
                    border: 2px solid #452829;
                }
            """)

            h_layout.addWidget(title_lbl)
            h_layout.addStretch()
            h_layout.addWidget(self.date_nav_box)
            h_layout.addStretch()
            h_layout.addWidget(self.custom_search)

            self.table_box = QFrame()
            self.table_box.setStyleSheet("QFrame { background-color: #f6f4f2; border-radius: 20px; border: 1px solid #e0e0e0; }")
            shadow = QtWidgets.QGraphicsDropShadowEffect()
            shadow.setBlurRadius(25); shadow.setColor(QColor(0,0,0,15)); shadow.setOffset(0,5)
            self.table_box.setGraphicsEffect(shadow)
            
            t_layout = QVBoxLayout(self.table_box)
            t_layout.setContentsMargins(15, 15, 15, 15)
            
            if hasattr(self, 'mainCalendarTable'):
                self.mainCalendarTable.setParent(self.table_box)
                t_layout.addWidget(self.mainCalendarTable)
                self.mainCalendarTable.show()
                self.mainCalendarTable.setStyleSheet("""
                    QTableWidget { background-color: transparent; border: none; gridline-color: #f7f7f7; color: #452829; }
                    QHeaderView::section { background-color: #fcfcfc; color: #452829; font-weight: bold; padding: 15px; border: none; border-bottom: 2px solid #eeeeee; }
                """)
                self.mainCalendarTable.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)

            self.verticalLayout_calendar.addWidget(self.header_card)
            self.verticalLayout_calendar.addWidget(self.table_box)
            self.verticalLayout_calendar.setStretch(1, 1)
            self.verticalLayout_calendar.setContentsMargins(20, 0, 20, 20)

            self.current_cal_date = date.today()
            self.custom_btn_prev.clicked.connect(lambda: self.show_calendar(self.current_cal_date - timedelta(days=1)))
            self.custom_btn_next.clicked.connect(lambda: self.show_calendar(self.current_cal_date + timedelta(days=1)))
            self.custom_search.textChanged.connect(self.filter_calendar_table)

        if hasattr(self, 'mainStack') and hasattr(self, 'page_empty'):
             self.mainStack.setCurrentWidget(self.page_empty)
        
        if hasattr(self, 'label_calendar_date'):
            self.label_calendar_date = self.custom_label_date
        
        # --- Apply Shadows for Premium Feel ---
        if hasattr(self, 'schedulingBottomCard'):
            shadow = QtWidgets.QGraphicsDropShadowEffect()
            shadow.setBlurRadius(30)
            shadow.setColor(QColor(0, 0, 0, 30))
            shadow.setOffset(0, 8)
            self.schedulingBottomCard.setGraphicsEffect(shadow)
            
        self.reset_sidebar_styles()

    def show_add_case(self):
        self.reset_sidebar_styles()
        self.add_case.setProperty("active", True)
        self.add_case.style().unpolish(self.add_case)
        self.add_case.style().polish(self.add_case)
        self.mainStack.setCurrentWidget(self.page_add_case)

        db = DataBase()
        db.cur.execute("""
            SELECT client_id, plaintiff_name, defendant_name, case_type
            FROM cms.client
            WHERE client_id NOT IN (SELECT client_id FROM cms.case_client)
        """)
        clients = db.cur.fetchall()
        db.close()

        table = self.addCaseTable
        table.setRowCount(0)
        table.verticalHeader().setVisible(False)
        table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.add_case_checkboxes = []

        if not clients:
            table.setRowCount(1)
            item = QtWidgets.QTableWidgetItem("لا توجد عملاء جدد لإضافتهم")
            item.setTextAlignment(Qt.AlignCenter)
            table.setItem(0, 0, item)
            table.setSpan(0, 0, 1, 4)
            self.btn_create_cases.setEnabled(False)
            return

        self.btn_create_cases.setEnabled(True)
        self._add_case_clients = clients

        for row, client in enumerate(clients):
            client_id, p_name, d_name, c_type = client
            table.insertRow(row)
            chk = QCheckBox()
            cell_widget = QWidget()
            layout = QHBoxLayout(cell_widget)
            layout.addWidget(chk)
            layout.setAlignment(Qt.AlignCenter)
            layout.setContentsMargins(0, 0, 0, 0)
            table.setCellWidget(row, 0, cell_widget)
            self.add_case_checkboxes.append(chk)
            for col, val in enumerate([p_name, d_name, c_type], start=1):
                item = QtWidgets.QTableWidgetItem(str(val))
                item.setTextAlignment(Qt.AlignCenter)
                table.setItem(row, col, item)
            
            table.setRowHeight(row, 50)
        
        header = table.horizontalHeader()
        header.setSectionResizeMode(QHeaderView.Stretch)

        try:
            self.btn_create_cases.clicked.disconnect()
        except:
            pass
        self.btn_create_cases.clicked.connect(self.create_cases_from_page)

    def create_cases_from_page(self):
        selected_rows = [i for i, cb in enumerate(self.add_case_checkboxes) if cb.isChecked()]
        if not selected_rows:
            QMessageBox.warning(self, "تنبيه", "اختر قضية واحدة على الأقل")
            return

        db = DataBase()
        seen_clients = set()
        for idx in selected_rows:
            client_id, p_name, d_name, c_type = self._add_case_clients[idx]
            if client_id in seen_clients:
                # duplicate selection somehow, skip
                continue
            seen_clients.add(client_id)
            # make sure we received defendant info
            if not d_name or d_name.strip() == "":
                QMessageBox.warning(self, "تنبيه", f"المدعى عليه للقضية {p_name} غير محدد في سجل الموكل؛ سيظهر كغير محدد في التقويم.")
            db.cur.execute("SELECT count(*) FROM cms.court_case")
            case_count = db.cur.fetchone()[0]
            case_number = f"{datetime.now().strftime('%Y')}/{case_count + 1}"
            db.cur.execute("""
                INSERT INTO cms.court_case (case_type, case_number, status, filing_date, year, description, created_by)
                VALUES (%s, %s, %s, %s, %s, %s, %s)
                RETURNING case_id
            """, (c_type, case_number, "مفتوحة", datetime.now().date(), datetime.now().year, "-", self.current_user_id))
            new_case_id = db.cur.fetchone()[0]

            # attach the client (which already contains both plaintiff/defendant names)
            try:
                db.cur.execute(
                    "INSERT INTO cms.case_client (case_id, client_id, role_in_case) VALUES (%s, %s, %s)",
                    (new_case_id, client_id, "Plaintiff")
                )
            except Exception as e:
                # ignore duplicate or other integrity errors gracefully
                if hasattr(e, 'pgcode') and e.pgcode == '23505':
                    print(f"warning: client {client_id} already linked to case {new_case_id}")
                else:
                    raise

            db.cur.execute("UPDATE cms.document SET case_id = %s WHERE client_id = %s AND case_id IS NULL",
                           (new_case_id, client_id))

            # Notify Judges
            db.cur.execute("SELECT user_id FROM cms.users WHERE role_id = 4")
            judges_to_notify = db.cur.fetchall()
            for judge_entry in judges_to_notify:
                case_date = datetime.now().strftime("%Y-%m-%d")
                notif_msg = f"تم إنشاء قضية جديدة: {c_type} برقم {case_number} بتاريخ {case_date}"
                db.cur.execute("INSERT INTO cms.notification (user_id, message) VALUES (%s, %s)", (judge_entry[0], notif_msg))

        db.conn.commit()
        db.close()
        QMessageBox.information(self, "نجاح", f"تم إنشاء {len(selected_rows)} قضية بنجاح ✅")
        self.show_add_case()

    def edit_selected_petition(self):
        selected_rows = [i for i, cb in enumerate(self.add_case_checkboxes) if cb.isChecked()]
        if len(selected_rows) != 1:
            QMessageBox.warning(self, "تنبيه", "يرجى اختيار قضية واحدة فقط للتعديل")
            return
        client_id = self._add_case_clients[selected_rows[0]][0]
        self.edit_petition(client_id)

    def delete_selected_petitions(self):
        selected_rows = [i for i, cb in enumerate(self.add_case_checkboxes) if cb.isChecked()]
        if not selected_rows:
            QMessageBox.warning(self, "تنبيه", "يرجى اختيار طلب واحد على الأقل للحذف")
            return
        
        reply = QMessageBox.question(self, "تأكيد الحذف", f"هل أنت متأكد من حذف {len(selected_rows)} طلب/عريضة نهائياً؟",
                                     QMessageBox.Yes | QMessageBox.No, QMessageBox.No)
        if reply == QMessageBox.Yes:
            try:
                db = DataBase()
                for idx in selected_rows:
                    client_id = self._add_case_clients[idx][0]
                    # Get documents
                    db.cur.execute("SELECT document_id, file_path FROM cms.document WHERE client_id = %s", (client_id,))
                    docs = db.cur.fetchall()
                    for doc_id, file_path in docs:
                        if file_path and os.path.exists(file_path):
                            try: os.remove(file_path)
                            except: pass
                        db.cur.execute("DELETE FROM cms.file_transfer WHERE document_id = %s", (doc_id,))
                        db.cur.execute("DELETE FROM cms.notification WHERE document_id = %s", (doc_id,))
                    db.cur.execute("DELETE FROM cms.document WHERE client_id = %s", (client_id,))
                    db.cur.execute("DELETE FROM cms.client WHERE client_id = %s", (client_id,))
                
                db.conn.commit()
                db.close()
                QMessageBox.information(self, "نجاح", "تم حذف الطلبات المحددة بنجاح ✅")
                self.show_add_case()
            except Exception as e:
                QMessageBox.critical(self, "خطأ", f"حدث خطأ أثناء الحذف: {str(e)}")

    def delete_petition(self, client_id):
        reply = QMessageBox.question(self, "تأكيد الحذف", "هل أنت متأكد من حذف هذه العريضة (طلب القضية) بجميع مستنداتها؟",
                                     QMessageBox.Yes | QMessageBox.No, QMessageBox.No)
        if reply == QMessageBox.Yes:
            try:
                db = DataBase()
                # 1. Get documents
                db.cur.execute("SELECT document_id, file_path FROM cms.document WHERE client_id = %s", (client_id,))
                docs = db.cur.fetchall()
                for doc_id, file_path in docs:
                    if file_path and os.path.exists(file_path):
                        try: os.remove(file_path)
                        except: pass
                    db.cur.execute("DELETE FROM cms.file_transfer WHERE document_id = %s", (doc_id,))
                    db.cur.execute("DELETE FROM cms.notification WHERE document_id = %s", (doc_id,))
                
                # 2. Delete documents
                db.cur.execute("DELETE FROM cms.document WHERE client_id = %s", (client_id,))
                
                # 3. Delete client
                db.cur.execute("DELETE FROM cms.client WHERE client_id = %s", (client_id,))
                
                db.conn.commit()
                db.close()
                QMessageBox.information(self, "نجاح", "تم حذف العريضة بنجاح ✅")
                self.show_add_case()
            except Exception as e:
                QMessageBox.critical(self, "خطأ", f"حدث خطأ أثناء الحذف: {str(e)}")

    def edit_petition(self, client_id):
        db = DataBase()
        db.cur.execute("""
            SELECT plaintiff_name, plaintiff_national_id, plaintiff_phone, plaintiff_address,
                   defendant_name, defendant_national_id, defendant_phone, defendant_address,
                   case_type 
            FROM cms.client WHERE client_id = %s
        """, (client_id,))
        client_data = db.cur.fetchone()
        db.close()
        
        if not client_data: return
        
        dialog = QtWidgets.QDialog(self)
        dialog.setWindowTitle("تعديل كافة بيانات الطلب")
        dialog.setFixedSize(550, 500)
        dialog.setLayoutDirection(Qt.RightToLeft)
        layout = QVBoxLayout(dialog)
        
        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setStyleSheet("QScrollArea { border: none; background: transparent; }")
        container = QWidget()
        form_layout = QtWidgets.QFormLayout(container)
        
        fields = {}
        labels = [
            ("اسم المدعي:", client_data[0]),
            ("رقم هوية المدعي:", client_data[1]),
            ("جوال المدعي:", client_data[2]),
            ("عنوان المدعي:", client_data[3]),
            ("اسم المدعى عليه:", client_data[4]),
            ("رقم هوية المدعى عليه:", client_data[5]),
            ("جوال المدعى عليه:", client_data[6]),
            ("عنوان المدعى عليه:", client_data[7]),
            ("نوع القضية:", client_data[8]),
        ]
        
        for i, (lbl_txt, val) in enumerate(labels):
            le = QtWidgets.QLineEdit(str(val) if val else "")
            le.setMinimumHeight(35)
            form_layout.addRow(lbl_txt, le)
            fields[i] = le
        
        scroll.setWidget(container)
        layout.addWidget(scroll)
        
        save_btn = QPushButton("حفظ كافة التغييرات")
        save_btn.setStyleSheet("background-color: #452829; color: white; padding: 12px; border-radius: 10px; font-weight: bold; font-size: 14px;")
        layout.addWidget(save_btn)
        
        def save_changes():
            try:
                db_save = DataBase()
                db_save.cur.execute("""
                    UPDATE cms.client 
                    SET plaintiff_name = %s, plaintiff_national_id = %s, plaintiff_phone = %s, plaintiff_address = %s,
                        defendant_name = %s, defendant_national_id = %s, defendant_phone = %s, defendant_address = %s,
                        case_type = %s
                    WHERE client_id = %s
                """, (
                    fields[0].text(), fields[1].text(), fields[2].text(), fields[3].text(),
                    fields[4].text(), fields[5].text(), fields[6].text(), fields[7].text(),
                    fields[8].text(), client_id
                ))
                db_save.conn.commit()
                db_save.close()
                QMessageBox.information(dialog, "نجاح", "تم تحديث كافة بيانات الطلب بنجاح ✅")
                dialog.accept()
                self.show_add_case()
            except Exception as e:
                QMessageBox.critical(dialog, "خطأ", f"حدث خطأ أثناء الحفظ: {str(e)}")
        
        save_btn.clicked.connect(save_changes)
        dialog.exec_()


    def get_hijri_date_string(self, date):
        try:
            jd = date.toordinal() + 1721425 + 1
            days_since_hijra = jd - 1948440
            cycles = days_since_hijra // 10631
            rem_days = days_since_hijra % 10631
            leap_years = [2, 5, 7, 10, 13, 16, 18, 21, 24, 26, 29]
            year_in_cycle = 0
            while True:
                year_in_cycle += 1
                is_leap = year_in_cycle in leap_years
                days_in_this_year = 355 if is_leap else 354
                if rem_days < days_in_this_year:
                    break
                rem_days -= days_in_this_year
            h_year = cycles * 30 + year_in_cycle
            h_month = 0
            for m in range(1, 13):
                h_month = m
                days_in_this_month = 30 if m % 2 != 0 else 29
                if m == 12 and (year_in_cycle in leap_years):
                    days_in_this_month = 30
                if rem_days < days_in_this_month:
                    break
                rem_days -= days_in_this_month
            h_day = rem_days + 1
            m_names = ["", "محرم", "صفر", "ربيع الأول", "ربيع الآخر", "جمادى الأولى", "جمادى الآخرة", 
                       "رجب", "شعبان", "رمضان", "شوال", "ذو القعدة", "ذو الحجة"]
            return f"{int(h_day)} {m_names[int(h_month)]} {int(h_year)} هـ"
        except Exception as e:
            print(f"Hijri Conversion Error: {e}")
            return "تاريخ غير متوفر"

    def update_badge(self):
        try:
            db = DataBase()
            db.cur.execute("""
                SELECT COUNT(*) FROM cms.notification 
                WHERE user_id = %s AND is_read = FALSE
            """, (self.current_user_id,))
            count = db.cur.fetchone()[0]
            db.close()

            if count > 0:
                self.badge_label.setText(str(count) if count < 10 else "9+")
                self.badge_label.show()
            else:
                self.badge_label.hide()
        except Exception as e:
            print(f"Error checking notifications: {e}")

    def show_notifications(self):
        db = DataBase()
        db.cur.execute("""
            SELECT notification_id, message, created_at, document_id 
            FROM cms.notification 
            WHERE user_id = %s AND is_read = FALSE
            ORDER BY created_at DESC
            LIMIT 10
        """, (self.current_user_id,))
        notifications = db.cur.fetchall()
        
        if notifications:
            ids = tuple([n[0] for n in notifications])
            if len(ids) == 1:
                db.cur.execute("DELETE FROM cms.notification WHERE notification_id = %s", (ids[0],))
            else:
                db.cur.execute("DELETE FROM cms.notification WHERE notification_id IN %s", (ids,))
            db.conn.commit()
        db.close()
        
        self.update_badge()

        menu = QMenu(self)
        menu.setMinimumWidth(350)
        menu.setStyleSheet("""
            QMenu {
                background-color: #452829;
                color: white;
                border: 1px solid #f3db93;
                border-radius: 10px;
                padding: 10px;
            }
        """)

        if not notifications:
            action = QWidgetAction(menu)
            lbl = QLabel("لا توجد إشعارات جديدة")
            lbl.setStyleSheet("color: #f3e8df; padding: 10px;")
            lbl.setAlignment(Qt.AlignCenter)
            action.setDefaultWidget(lbl)
            menu.addAction(action)
        else:
            for notif_id, msg, created_at, doc_id in notifications:
                time_str = created_at.strftime("%I:%M %p")
                item_widget = QWidget()
                item_widget.setStyleSheet("background-color: transparent;")
                item_layout = QHBoxLayout(item_widget)
                item_layout.setContentsMargins(10, 5, 10, 5)
                item_layout.setDirection(QHBoxLayout.RightToLeft)
                
                msg_label = QLabel(msg)
                msg_label.setStyleSheet("color: white; font-weight: bold; font-family: 'Alyamama'; font-size: 14px;")
                msg_label.setWordWrap(True)
                msg_label.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
                
                time_label = QLabel(time_str)
                time_label.setStyleSheet("color: #f3db93; font-size: 12px;")
                time_label.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)

                sep = QLabel("|")
                sep.setStyleSheet("color: rgba(255,255,255,0.2); font-weight: bold; font-size: 14px;")

                item_layout.addWidget(msg_label, 1)
                item_layout.addWidget(sep, 0)
                item_layout.addWidget(time_label, 0)
                
                action = QWidgetAction(menu)
                action.setDefaultWidget(item_widget)
                
                if doc_id:
                     action.triggered.connect(lambda checked=False, d=doc_id: self.handle_notification_click(d))
                
                menu.addAction(action)
                menu.addSeparator()

        menu.exec_(self.notification.mapToGlobal(QPoint(0, self.notification.height())))

    def handle_notification_click(self, document_id):
        self.show_documents(highlight_id=document_id)

    def reset_sidebar_styles(self):
        buttons = [self.add_case, self.docments, self.master_record, self.btn_scheduling, self.calendar]
        for btn in buttons:
            btn.setFocusPolicy(Qt.NoFocus)
            btn.setProperty("active", False)
            btn.style().unpolish(btn)
            btn.style().polish(btn)

    def show_documents(self, highlight_id=None):
        self.reset_sidebar_styles()
        self.docments.setProperty("active", True)
        self.docments.style().unpolish(self.docments)
        self.docments.style().polish(self.docments)
        if hasattr(self, 'mainStack'):
            self.mainStack.setCurrentWidget(self.page_documents)

        try:
            db_clear = DataBase()
            db_clear.cur.execute("DELETE FROM cms.notification WHERE user_id = %s", (self.current_user_id,))
            db_clear.conn.commit()
            db_clear.close()
            self.update_badge()
        except Exception as e:
            print(f"Error clearing notifications: {e}")

        if hasattr(self, 'files_grid'):
            self.selected_documents.clear()
            self.doc_checkboxes.clear()
            if hasattr(self, 'check_all_docs'):
                self.check_all_docs.blockSignals(True)
                self.check_all_docs.setChecked(False)
                self.check_all_docs.blockSignals(False)
                
            while self.files_grid.count():
                item = self.files_grid.takeAt(0)
                widget = item.widget()
                if widget:
                    widget.deleteLater()

        db = DataBase()
        db.cur.execute("""
            SELECT d.file_path, d.document_id, n.created_at
            FROM cms.file_transfer ft
            JOIN cms.document d ON ft.document_id = d.document_id
            LEFT JOIN cms.notification n ON d.document_id = n.document_id AND n.user_id = ft.receiver_id
            WHERE ft.receiver_id = %s
            ORDER BY ft.transfer_date DESC
        """, (self.current_user_id,))
        files = db.cur.fetchall()
        db.close()

        row_idx = 0
        for (file_path, doc_id, created_at) in files:
            row_widget = QWidget()
            row_widget.setFixedHeight(80)
            
            normal_style = """
                QWidget {
                    background-color: white;
                    border-bottom: 1px solid #e0e0e0;
                }
                QWidget:hover {
                    background-color: #f9f9f9;
                }
            """
            
            if highlight_id and doc_id == highlight_id:
                highlight_style = normal_style + "QWidget { background-color: #fff8e1; }"
                row_widget.setStyleSheet(highlight_style)
                QTimer.singleShot(3000, lambda w=row_widget: w.setStyleSheet(normal_style))
            else:
                row_widget.setStyleSheet(normal_style)

            layout = QHBoxLayout(row_widget)
            layout.setContentsMargins(20, 10, 20, 10)
            layout.setSpacing(15)

            checkbox = QCheckBox()
            checkbox.setCursor(Qt.PointingHandCursor)
            checkbox.setStyleSheet("QCheckBox::indicator { width: 20px; height: 20px; }")
            checkbox.setProperty("doc_id", doc_id)
            checkbox.stateChanged.connect(lambda state, d=doc_id: self.toggle_doc_selection(d, state))
            self.doc_checkboxes.append(checkbox)
            layout.addWidget(checkbox)
            
            icon = QLabel("📄")
            icon.setStyleSheet("font-size: 30px; background: transparent; border: none;")
            
            file_name = os.path.basename(file_path)
            name_label = QLabel(file_name)
            name_label.setStyleSheet("font-size: 16px; font-weight: bold; color: #333; background: transparent; border: none;")
            name_label.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
            
            time_str = created_at.strftime("%I:%M %p") if created_at else ""
            time_label = QLabel(time_str)
            time_label.setStyleSheet("color: #777; font-size: 14px; background: transparent; border: none; font-family: 'Alyamama';")
            time_label.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
            
            spacer = QWidget()
            spacer.setSizePolicy(QtWidgets.QSizePolicy.Expanding, QtWidgets.QSizePolicy.Preferred)
            spacer.setStyleSheet("background: transparent; border: none;")

            btn_extract = QPushButton("استخراج إعلان الخصوم")
            btn_extract.setCursor(Qt.PointingHandCursor)
            btn_extract.setFocusPolicy(Qt.NoFocus)
            btn_extract.setMinimumHeight(40)
            btn_extract.setStyleSheet("""
                QPushButton {
                    background-color: #452829; 
                    color: white; 
                    border-radius: 5px; 
                    padding: 5px 15px;
                    border: none;
                }
                QPushButton:hover {
                    background-color: #f3db93;
                    color: black;
                }
            """)
            btn_extract.clicked.connect(lambda checked=False, d=doc_id: self.extract_notification_file(d))

            btn_open = QPushButton("فتح")
            btn_open.setCursor(Qt.PointingHandCursor)
            btn_open.setFocusPolicy(Qt.NoFocus)
            btn_open.setMinimumHeight(40)
            btn_open.setStyleSheet("""
                QPushButton {
                    background-color: #452829; 
                    color: white; 
                    border-radius: 5px; 
                    padding: 5px 20px;
                    border: none;
                }
                QPushButton:hover {
                    background-color: #f3db93;
                    color: black;
                }
            """)
            btn_open.clicked.connect(lambda checked=False, p=file_path: self.open_file(p))

            layout.addWidget(icon)
            layout.addWidget(name_label)
            layout.addWidget(time_label)
            layout.addWidget(spacer)
            layout.addWidget(btn_extract)
            layout.addWidget(btn_open)

            self.files_grid.addWidget(row_widget, row_idx, 0)
            row_idx += 1

    def toggle_doc_selection(self, doc_id, state):
        if state == Qt.Checked:
            self.selected_documents.add(doc_id)
        else:
            self.selected_documents.discard(doc_id)

    def select_all_documents(self, state):
        is_checked = (state == Qt.Checked)
        for cb in self.doc_checkboxes:
            cb.setChecked(is_checked)

    def delete_selected_documents(self):
        if not self.selected_documents:
            QMessageBox.information(self, "تنبيه", "يرجى اختيار المستندات المراد حذفها أولاً.")
            return

        confirm = QMessageBox.question(
            self, "تأكيد الحذف",
            f"هل أنت متأكد من حذف {len(self.selected_documents)} من المستندات المختارة؟\nهذا الإجراء سيحذف السجلات من النظام بشكل نهائي.",
            QMessageBox.Yes | QMessageBox.No
        )
        
        if confirm == QMessageBox.Yes:
            try:
                db = DataBase()
                for doc_id in self.selected_documents:
                    db.cur.execute("DELETE FROM cms.notification WHERE document_id = %s", (doc_id,))
                    db.cur.execute("DELETE FROM cms.file_transfer WHERE document_id = %s", (doc_id,))
                    db.cur.execute("DELETE FROM cms.document WHERE document_id = %s", (doc_id,))
                db.conn.commit()
                db.close()
                QMessageBox.information(self, "نجاح", "تم حذف المستندات المختارة بنجاح.")
                self.show_documents()
            except Exception as e:
                QMessageBox.critical(self, "خطأ", f"حدث خطأ أثناء الحذف:\n{str(e)}")

    def extract_notification_file(self, doc_id):
        try:
            db = DataBase()
            db.cur.execute("""
                SELECT 
                    COALESCE(c1.plaintiff_name, c2.plaintiff_name) as plaintiff_name,
                    COALESCE(c1.plaintiff_address, c2.plaintiff_address) as plaintiff_address,
                    COALESCE(c1.defendant_name, c2.defendant_name) as defendant_name,
                    COALESCE(c1.defendant_address, c2.defendant_address) as defendant_address,
                    d.document_type,
                    d.upload_date,
                    j.full_name as judge_name,
                    ct.case_number,
                    s.session_date,
                    s.session_time
                FROM cms.document d
                LEFT JOIN cms.users u ON d.uploaded_by = u.user_id
                LEFT JOIN cms.client c1 ON d.client_id = c1.client_id
                LEFT JOIN cms.case_client cc ON (cc.case_id = d.case_id OR (d.case_id IS NULL AND cc.client_id = d.client_id))
                LEFT JOIN cms.client c2 ON cc.client_id = c2.client_id
                LEFT JOIN cms.court_case ct ON cc.case_id = ct.case_id
                LEFT JOIN cms.session s ON cc.case_id = s.case_id
                LEFT JOIN cms.users j ON s.judge_id = j.user_id
                WHERE d.document_id = %s
                ORDER BY s.session_id DESC NULLS LAST, ct.filing_date DESC LIMIT 1
            """, (doc_id,))
            result = db.cur.fetchone()
            db.close()

            if not result:
                QMessageBox.warning(self, "خطأ", "لم يتم العثور على بيانات الموكل للمستند المحدد")
                return
            
            plaintiff_name, plaintiff_address, defendant_name, defendant_address, doc_type, upload_date, judge_name, case_number, session_date_val, session_time_val = result

            if not plaintiff_name:
                QMessageBox.warning(self, "تنبيه", "بيانات الموكل (المدعي) ناقصة أو غير مرتبطة بشكل صحيح بهذا المستند.")
                return

            plaintiff_address = plaintiff_address if plaintiff_address else ""
            defendant_name = defendant_name if defendant_name else ""
            defendant_address = defendant_address if defendant_address else ""
            doc_type = doc_type if doc_type else "-"
            case_number = case_number if case_number else ""

            # تحويل تاريخ ووقت الجلسة
            arabic_days = {
                0: "الاثنين", 1: "الثلاثاء", 2: "الأربعاء",
                3: "الخميس", 4: "الجمعة", 5: "السبت", 6: "الأحد"
            }

            if session_date_val:
                session_day_str = arabic_days.get(session_date_val.weekday(), "")
                session_date_str = session_date_val.strftime("%Y/%m/%d")
            else:
                session_day_str = ""
                session_date_str = ""

            if session_time_val:
                if hasattr(session_time_val, 'strftime'):
                    session_time_str = session_time_val.strftime("%H:%M")
                else:
                    session_time_str = str(session_time_val)[:5]
            else:
                session_time_str = ""

            template_path = os.path.join("Template files", "إعلان خصوم.docx")
            if not os.path.exists(template_path):
                QMessageBox.warning(self, "خطأ", f"قالب إعلان الخصوم غير موجود:\n{template_path}")
                return
            
            safe_name = "".join([c for c in plaintiff_name if c.isalnum() or c in (' ', '_')]).rstrip()
            final_filename = f"إعلان خصوم - {safe_name}.docx"
            final_file_path = os.path.join(self.db.files_path, final_filename)
            shutil.copy2(template_path, final_file_path)
            
            doc = Document(final_file_path)
            
            p_addr = plaintiff_address if plaintiff_address else ""
            p_from = p_addr.split('-')[0].strip() if '-' in p_addr else p_addr
            p_res = p_addr.split('-')[1].strip() if '-' in p_addr else "-"

            d_addr = defendant_address if defendant_address else ""
            d_from = d_addr.split('-')[0].strip() if '-' in d_addr else d_addr
            d_res = d_addr.split('-')[1].strip() if '-' in d_addr else "-"

            placeholders = {
                "{PLAINTIFF_NAME}": plaintiff_name or "",
                "{PLAINTIFF_FROM}": p_from,
                "{PLAINTIFF_RESIDENT}": p_res,
                "{DEFENDANT_NAME}": defendant_name or "",
                "{DEFENDANT_FROM}": d_from,
                "{DEFENDANT_RESIDENT}": d_res,
                "{CURRENT_CASE_TYPE}": f" {doc_type}" if doc_type else "",
                "{ENTRY_DATE}": upload_date.strftime("%Y/%m/%d") if upload_date else datetime.now().strftime("%Y/%m/%d"),
                "{JUDGE_NAME}": judge_name or "",
                "{CASE_NUMBER}": case_number,
                "{SESSION_DAY}": session_day_str,
                "{SESSION_DATE}": session_date_str,
                "{SESSION_TIME}": session_time_str,
            }

            from doc_helpers import safe_replace_in_doc
            safe_replace_in_doc(doc, placeholders)
            doc.save(final_file_path)
            
            # Save extracted notification document back to database so Judge can view it
            if result:
                db_save = DataBase()
                db_save.cur.execute("SELECT case_id, client_id FROM cms.document WHERE document_id = %s", (doc_id,))
                doc_meta = db_save.cur.fetchone()
                if doc_meta:
                    orig_case_id, orig_client_id = doc_meta
                    
                    # Check if notification already exists for this case
                    db_save.cur.execute(
                        "SELECT document_id FROM cms.document WHERE case_id = %s AND document_type = %s",
                        (orig_case_id, "إعلان خصوم")
                    )
                    existing = db_save.cur.fetchone()
                    
                    if existing:
                        # Update existing one
                        db_save.cur.execute(
                            "UPDATE cms.document SET file_path = %s, upload_date = CURRENT_TIMESTAMP WHERE document_id = %s",
                            (final_file_path, existing[0])
                        )
                    else:
                        # Insert new one
                        db_save.cur.execute(
                            "INSERT INTO cms.document (document_type, file_path, uploaded_by, case_id, client_id) VALUES (%s, %s, %s, %s, %s)",
                            ("إعلان خصوم", final_file_path, self.current_user_id, orig_case_id, orig_client_id)
                        )
                    db_save.conn.commit()
                db_save.close()

            QMessageBox.information(self, "نجاح", f"تم استخراج وتعديل إعلان الخصوم بنجاح!\n\nالملف: {final_filename}")
            
            try:
                os.startfile(final_file_path)
            except Exception as e:
                print(f"Could not open file: {e}")
                
        except Exception as e:
            QMessageBox.critical(self, "خطأ", f"حدث خطأ أثناء استخراج إعلان الخصوم:\n{str(e)}")
            
    def open_file(self, file_path):
        try:
            # Normalize path
            file_path = file_path.replace('/', '\\')
            
            # 1. Try the stored path directly
            if os.path.exists(file_path):
                os.startfile(os.path.abspath(file_path))
                return
            
            # 2. Try relative to the configured files_path (Crucial for network/other machines)
            filename = os.path.basename(file_path)
            net_path = os.path.join(self.db.files_path, filename)
            
            if os.path.exists(net_path):
                os.startfile(net_path)
            else:
                QMessageBox.warning(self, "خطأ في فتح الملف", 
                    f"لم يتم العثور على الملف في المسار المحلي أو في مسار الشبكة:\n\n"
                    f"المسار المحاول: {net_path}\n\n"
                    "تأكد من أن جهاز السيرفر (MoAlshanti) متاح وأن مجلد files تمت مشاركته.")
        except Exception as e:
            QMessageBox.warning(self, "خطأ", f"تعذر فتح الملف: {e}")
            
    def log_out(self):
        if self.main_shell:
            self.main_shell.switch_to_login()
        else:
            self.close()

    def show_calendar(self, selected_qdate=None):
        db = DataBase()
        self.reset_sidebar_styles()
        self.calendar.setProperty("active", True)
        self.calendar.style().unpolish(self.calendar)
        self.calendar.style().polish(self.calendar)
        
        self.mainStack.setCurrentWidget(self.page_calendar)
        
        if selected_qdate:
             if hasattr(selected_qdate, 'toPyDate'):
                 selected_date = selected_qdate.toPyDate()
             elif isinstance(selected_qdate, (date, datetime)):
                 selected_date = selected_qdate
             else:
                 selected_date = date.today()
        else:
             selected_date = date.today()

        self.current_cal_date = selected_date

        hijri_str = self.get_hijri_date_string(selected_date)
        greg_str = selected_date.strftime("%d %B %Y")
        combined_date = f"{hijri_str}  |  {greg_str}"
        
        if hasattr(self, 'label_calendar_date'):
             self.label_calendar_date.setText(combined_date)
             self.label_calendar_date.setStyleSheet("color: #452829; font-size: 18px; font-family: 'Alyamama';")
        
        db.cur.execute("""
            SELECT user_id, full_name FROM cms.users WHERE role_id = 4 ORDER BY user_id
        """)
        judges_data = db.cur.fetchall()
        
        judge_names = ["ساعة"] + [f"القاضي {judge[1]} " for judge in judges_data]
        j_id = [(j[0],) for j in judges_data] 
        
        table = self.mainCalendarTable
        table.setColumnCount(len(judge_names))
        table.setHorizontalHeaderLabels(judge_names)
        
        hours_list = [f"{h:02d}:00" for h in range(8, 17)]
        table.setRowCount(len(hours_list))
        table.verticalHeader().setVisible(False)
        
        table.horizontalHeader().setStyleSheet("""
            QHeaderView::section { 
                background-color: #fcfcfc; 
                font-weight: bold; 
                border: 1px solid #e0e0e0; 
                color: #452829;
                padding: 10px;
                font-family: 'Alyamama';
            }
        """)
        table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        table.setEditTriggers(QtWidgets.QAbstractItemView.NoEditTriggers)
        table.clearContents()
        
        def add_session_block(row, col, text, color_type="maroon"):
            colors = {
                "maroon": "#452829",
                "gold": "#b08d57", 
                "beige": "#ebe3d5"
            }
            text_colors = {
                "maroon": "white",
                "gold": "white",
                "beige": "#452829"
            }
            
            bg_color = colors.get(color_type, "#452829")
            fg_color = text_colors.get(color_type, "white")
            
            container = QWidget()
            container_layout = QVBoxLayout(container)
            container_layout.setContentsMargins(4, 4, 4, 4)
            
            card = QFrame()
            card.setStyleSheet(f"background-color: {bg_color}; border-radius: 8px; border: none;")
            card_layout = QVBoxLayout(card)
            card_layout.setContentsMargins(8, 8, 8, 8)
            
            lbl = QLabel(text)
            lbl.setAlignment(Qt.AlignCenter)
            lbl.setWordWrap(True)
            lbl.setStyleSheet(f"color: {fg_color}; font-weight: bold; font-family: 'Alyamama'; font-size: 11px; background: transparent; border: none;")
            
            card_layout.addWidget(lbl)
            container_layout.addWidget(card)
            table.setCellWidget(row, col, container)

        for i, h in enumerate(hours_list):
            item = QtWidgets.QTableWidgetItem(h)
            item.setTextAlignment(Qt.AlignCenter)
            item.setForeground(QColor("#452829"))
            font = item.font()
            font.setBold(True)
            item.setFont(font)
            table.setItem(i, 0, item)

        today_str = selected_date.strftime("%Y-%m-%d")
        judge_id_to_col = {jid[0]: idx + 1 for idx, jid in enumerate(j_id)}
        
        db.cur.execute("""
            SELECT s.session_time, s.case_id, s.judge_id, c.case_number, c.case_type
            FROM cms.session s
            JOIN cms.court_case c ON s.case_id = c.case_id
            WHERE s.session_date = %s AND s.status = 'Scheduled'
        """, (today_str,))
        
        all_sessions = db.cur.fetchall()
        
        for s_time_val, case_id_val, judge_id_val, case_num, case_type in all_sessions:
            if judge_id_val not in judge_id_to_col: continue
            
            col_idx = judge_id_to_col[judge_id_val]
            
            if hasattr(s_time_val, 'strftime'):
                t_str = s_time_val.strftime("%H:00")
            elif hasattr(s_time_val, 'hour'):
                t_str = f"{s_time_val.hour:02d}:00"
            else:
                t_str = str(s_time_val)[:2] + ":00"
            
            if t_str in hours_list:
                row_idx = hours_list.index(t_str)
                color_choice = ["maroon", "gold", "beige"][case_id_val % 3]
                add_session_block(row_idx, col_idx, f"{case_type}\n{case_num}", color_choice)

        for i in range(len(hours_list)):
            table.setRowHeight(i, 70)
        db.close()

    def filter_calendar_table(self, text):
        table = self.mainCalendarTable
        search_term = text.lower()
        
        for row in range(table.rowCount()):
            for col in range(table.columnCount()):
                widget = table.cellWidget(row, col)
                if widget:
                    children = widget.findChildren(QLabel)
                    found = False
                    for child in children:
                        if search_term in child.text().lower():
                            found = True
                            break
                    
                    if not search_term:
                         widget.setHidden(False)
                    elif found:
                         widget.setHidden(False)
                    else:
                         widget.setHidden(True)

    def filter_master_record(self, text):
        table = self.masterRecordTable
        for row in range(table.rowCount()):
            match = False
            for col in range(table.columnCount()):
                item = table.item(row, col)
                if item and text.lower() in item.text().lower():
                    match = True
                    break
            table.setRowHidden(row, not match)

    def filter_scheduling(self, text):
        table = self.schedulingTable
        for row in range(table.rowCount()):
            match = False
            for col in range(1, table.columnCount()):
                item = table.item(row, col)
                if item and text.lower() in item.text().lower():
                    match = True
                    break
            table.setRowHidden(row, not match)

    def show_master_record(self):
        self.reset_sidebar_styles()
        self.master_record.setProperty("active", True)
        self.master_record.style().unpolish(self.master_record)
        self.master_record.style().polish(self.master_record)
        
        db = DataBase()
        db.cur.execute("""
            SELECT ct.case_number, c.plaintiff_name, c.defendant_name, ct.case_type, ct.filing_date, ct.status, ct.case_id
            FROM cms.case_client cc
            JOIN cms.client c ON cc.client_id = c.client_id
            JOIN cms.court_case ct ON cc.case_id = ct.case_id
            ORDER BY ct.filing_date ASC, CAST(SPLIT_PART(ct.case_number, '/', 2) AS INTEGER) ASC
        """)
        records = db.cur.fetchall()
        db.close()

        table = self.masterRecordTable
        table.verticalHeader().setVisible(False)
        table.setRowCount(0)
        table.setRowCount(len(records))
        table.verticalHeader().setDefaultSectionSize(50) # Increase row height
        
        header = table.horizontalHeader()
        header.setSectionResizeMode(QHeaderView.Stretch)
        
        for row_idx, row_data in enumerate(records):
            table.setRowHeight(row_idx, 50)
            for col_idx, value in enumerate(row_data[:6]):
                item = QtWidgets.QTableWidgetItem(str(value))
                item.setTextAlignment(Qt.AlignCenter)
                if col_idx == 5:
                    status_text = str(value)
                    if status_text in ["جديد", "مفتوحة"]:
                        item.setForeground(QColor("#2ECC71")) # Green
                    elif status_text in ["مغلق", "منتهية"]:
                        item.setForeground(QColor("#E74C3C")) # Red
                    font = item.font()
                    font.setBold(True)
                    item.setFont(font)
                table.setItem(row_idx, col_idx, item)

        self.searchMasterRecord.clear()
        self.mainStack.setCurrentWidget(self.page_master_record)


    def delete_case(self, case_id):
        reply = QMessageBox.question(self, "تأكيد الحذف", "هل أنت متأكد من حذف هذه القضية بجميع بياناتها؟",
                                     QMessageBox.Yes | QMessageBox.No, QMessageBox.No)
        if reply == QMessageBox.Yes:
            try:
                db = DataBase()
                # 1. Delete notifications related to this case? 
                # (Need to check if notifications are linked to case_id, usually they are linked via message or document_id)
                
                # 2. Get document IDs related to this case
                db.cur.execute("SELECT document_id, file_path FROM cms.document WHERE case_id = %s", (case_id,))
                docs = db.cur.fetchall()
                
                for doc_id, file_path in docs:
                    # Delete physical file
                    if file_path and os.path.exists(file_path):
                        try: os.remove(file_path)
                        except: pass
                    
                    # Delete file transfers
                    db.cur.execute("DELETE FROM cms.file_transfer WHERE document_id = %s", (doc_id,))
                    # Delete notifications
                    db.cur.execute("DELETE FROM cms.notification WHERE document_id = %s", (doc_id,))
                
                # 3. Delete documents
                db.cur.execute("DELETE FROM cms.document WHERE case_id = %s", (case_id,))
                
                # 4. Delete sessions
                db.cur.execute("DELETE FROM cms.session WHERE case_id = %s", (case_id,))
                
                # 5. Delete case-client links
                db.cur.execute("DELETE FROM cms.case_client WHERE case_id = %s", (case_id,))
                
                # 6. Delete the case itself
                db.cur.execute("DELETE FROM cms.court_case WHERE case_id = %s", (case_id,))
                
                db.conn.commit()
                db.close()
                QMessageBox.information(self, "نجاح", "تم حذف القضية بنجاح ✅")
                self.show_master_record()
            except Exception as e:
                QMessageBox.critical(self, "خطأ", f"حدث خطأ أثناء الحذف: {str(e)}")

    def edit_case(self, case_id):
        db = DataBase()
        db.cur.execute("""
            SELECT case_number, case_type, status, filing_date 
            FROM cms.court_case WHERE case_id = %s
        """, (case_id,))
        case_data = db.cur.fetchone()
        db.close()
        
        if not case_data: return
        
        dialog = QtWidgets.QDialog(self)
        dialog.setWindowTitle("تعديل بيانات القضية")
        dialog.setFixedSize(400, 300)
        dialog.setLayoutDirection(Qt.RightToLeft)
        layout = QVBoxLayout(dialog)
        
        form_layout = QtWidgets.QFormLayout()
        
        num_edit = QtWidgets.QLineEdit(case_data[0])
        type_edit = QtWidgets.QLineEdit(case_data[1])
        status_combo = QtWidgets.QComboBox()
        status_combo.addItems(["مفتوحة", "مغلق", "مؤجل"])
        status_combo.setCurrentText(case_data[2])
        
        form_layout.addRow("رقم القضية:", num_edit)
        form_layout.addRow("نوع القضية:", type_edit)
        form_layout.addRow("حالة القضية:", status_combo)
        
        layout.addLayout(form_layout)
        
        save_btn = QPushButton("حفظ التغييرات")
        save_btn.setStyleSheet("background-color: #452829; color: white; padding: 10px; border-radius: 5px; font-weight: bold;")
        layout.addWidget(save_btn)
        
        def save_changes():
            try:
                db_save = DataBase()
                db_save.cur.execute("""
                    UPDATE cms.court_case 
                    SET case_number = %s, case_type = %s, status = %s
                    WHERE case_id = %s
                """, (num_edit.text(), type_edit.text(), status_combo.currentText(), case_id))
                db_save.conn.commit()
                db_save.close()
                QMessageBox.information(dialog, "نجاح", "تم تحديث بيانات القضية بنجاح ✅")
                dialog.accept()
                self.show_master_record()
            except Exception as e:
                QMessageBox.critical(dialog, "خطأ", f"حدث خطأ أثناء الحفظ: {str(e)}")
        
        save_btn.clicked.connect(save_changes)
        dialog.exec_()
        # Sorting is already handled correctly by the SQL query above (numeric order)
        # Do NOT call table.sortItems here as it sorts by string, causing "10" < "2" bug

        self.searchMasterRecord.clear()
        self.mainStack.setCurrentWidget(self.page_master_record)

    def show_scheduling(self):
        self.reset_sidebar_styles()
        self.btn_scheduling.setProperty("active", True)
        self.btn_scheduling.style().unpolish(self.btn_scheduling)
        self.btn_scheduling.style().polish(self.btn_scheduling)
        self.mainStack.setCurrentWidget(self.page_scheduling)
        
        db = DataBase()
        db.cur.execute("""
            SELECT cc.case_id, ct.case_number, c.plaintiff_name, c.defendant_name, ct.case_type
            FROM cms.case_client cc
            JOIN cms.client c ON cc.client_id = c.client_id
            JOIN cms.court_case ct ON cc.case_id = ct.case_id
            WHERE cc.case_id NOT IN (
                SELECT case_id FROM cms.session WHERE status = 'Scheduled'
            )
            ORDER BY ct.filing_date DESC
        """)
        records = db.cur.fetchall()
        db.close()
        
        self.searchScheduling.clear()
        table = self.schedulingTable
        table.setRowCount(0)
        table.verticalHeader().setVisible(False)
        table.setRowCount(len(records))
        table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        
        self.scheduling_checkboxes = [] 
        table.horizontalHeader().setSectionResizeMode(0, QHeaderView.Fixed)
        table.setColumnWidth(0, 120) # Significantly increased width for "اختيار الجلسة"
        table.verticalHeader().setDefaultSectionSize(55)
        
        for row, data in enumerate(records):
            # data = (case_id, case_number, plaintiff, defendant, case_type)
            table.setRowHeight(row, 55)
            chk = QCheckBox()
            chk.setFixedSize(20, 20)
            
            cell_widget = QWidget()
            layout = QHBoxLayout(cell_widget)
            layout.addWidget(chk)
            layout.setAlignment(Qt.AlignCenter)
            layout.setContentsMargins(0, 0, 0, 0)
            table.setCellWidget(row, 0, cell_widget)
            # store case_id separately, but don't show it to user
            self.scheduling_checkboxes.append((chk, data[0]))
            
            # display case number, plaintiff, defendant, type
            for col, val in enumerate(data[1:], start=1):
                item = QtWidgets.QTableWidgetItem(str(val))
                item.setTextAlignment(Qt.AlignCenter)
                table.setItem(row, col, item)

        self.judgeComboBox.clear()
        self.judgeComboBox.addItem("اختر القاضي")
        db = DataBase()
        db.cur.execute("SELECT user_id, full_name FROM cms.users WHERE role_id = 4")
        judges = db.cur.fetchall()
        db.close()
        for judge in judges:
            self.judgeComboBox.addItem(judge[1], judge[0])

        if hasattr(self, 'sessionTimeInput'):
            self.sessionTimeInput.setMinimumTime(QTime(8, 0))
            self.sessionTimeInput.setMaximumTime(QTime(14, 59))
            self.sessionTimeInput.setTime(QTime(8, 0))
            # Force focus on Hour section so arrows work without manual selection
            self.sessionTimeInput.setFocus()
            self.sessionTimeInput.setCurrentSection(QtWidgets.QDateTimeEdit.HourSection)

        if hasattr(self, 'sessionDateInput'):
            current = date.today()
            self.sessionDateInput.setMinimumDate(current) # Prevent past dates
            
            # If today is weekend, move to next Sunday
            # weekday(): 0=Mon, 4=Fri, 5=Sat, 6=Sun
            if current.weekday() == 4: # Friday
                self.sessionDateInput.setDate(current + timedelta(days=2))
            elif current.weekday() == 5: # Saturday
                self.sessionDateInput.setDate(current + timedelta(days=1))
            else:
                self.sessionDateInput.setDate(current)

    def save_session(self):
        selected_case_id = None
        if not hasattr(self, 'scheduling_checkboxes'): return

        for chk, case_id in self.scheduling_checkboxes:
            if chk.isChecked():
                selected_case_id = case_id
                break
        
        if not selected_case_id:
            QMessageBox.warning(self, "تنبيه", "الرجاء اختيار قضية")
            return
            
        judge_id = self.judgeComboBox.currentData()
        if not judge_id:
            QMessageBox.warning(self, "تنبيه", "الرجاء اختيار القاضي")
            return

        val_time = self.sessionTimeInput.time()
        if val_time.hour() < 8 or val_time.hour() > 14:
             QMessageBox.warning(self, "تنبيه", "وقت الجلسة يجب أن يكون بين 8 صباحاً و 2 ظهراً (14:00) ليظهر في الجدول بشكل صحيح.")
             return

        q_date = self.sessionDateInput.date()
        session_date = date(q_date.year(), q_date.month(), q_date.day())
        session_time = time(val_time.hour(), val_time.minute())
        
        # Weekend Check (Friday=4, Saturday=5)
        if session_date.weekday() in [4, 5]:
             QMessageBox.warning(self, "تنبيه", "لا يمكن جدولة جلسات في أيام الإجازة (الجمعة والسبت).")
             return

        # Past Date Check (Just in case)
        if session_date < date.today():
             QMessageBox.warning(self, "تنبيه", "لا يمكن جدولة جلسات في تاريخ سابق.")
             return
        
        QtWidgets.QApplication.setOverrideCursor(Qt.WaitCursor)
        db = DataBase()
        try:
            db.cur.execute("""
                SELECT count(*) FROM cms.session 
                WHERE judge_id = %s AND session_date = %s AND session_time = %s AND status = 'Scheduled'
            """, (judge_id, session_date, session_time))
            if db.cur.fetchone()[0] > 0:
                QtWidgets.QApplication.restoreOverrideCursor()
                QMessageBox.warning(self, "تنبيه", "يوجد جلسة أخرى لهذا القاضي في نفس الموعد!")
                return

            db.cur.execute("""
                INSERT INTO cms.session (session_date, session_time, status, case_id, judge_id)
                VALUES (%s, %s, %s, %s, %s)
            """, (session_date, session_time, 'Scheduled', selected_case_id, judge_id))
            db.conn.commit()
            db.close()

            db_case = DataBase()
            db_case.cur.execute("SELECT case_number FROM cms.court_case WHERE case_id = %s", (selected_case_id,))
            case_number_val = db_case.cur.fetchone()[0]

            db_case.cur.execute("""
                SELECT d.file_path 
                FROM cms.document d
                INNER JOIN cms.case_client cc ON d.client_id = cc.client_id
                WHERE cc.case_id = %s
            """, (selected_case_id,))
            documents = db_case.cur.fetchall()
            db_case.close()
            
            if not documents:
                print(f"Warning: No documents found for case_id {selected_case_id}")
                # Don't return, maybe the user wants to know session was saved anyway
            
            arabic_days = {
                0: "الاثنين", 1: "الثلاثاء", 2: "الأربعاء",
                3: "الخميس", 4: "الجمعة", 5: "السبت", 6: "الأحد"
            }
            session_day_str = arabic_days.get(session_date.weekday(), "")
            date_str = session_date.strftime("%Y-%m-%d")
            time_str = session_time.strftime("%H:%M")
            
            placeholders = {          
                "{CASE_NUMBER}": str(case_number_val).translate(str.maketrans("٠١٢٣٤٥٦٧٨٩", "0123456789")),
                "{SESSION_DATE}": f"{session_date.year}/{session_date.month:02d}/{session_date.day:02d}",
                "{SESSION_DAY}": session_day_str,
                "{SESSION_TIME}": f"{session_time.hour:02d}:{session_time.minute:02d}",
                "{ENTRY_DATE}": datetime.now().strftime("%Y/%m/%d"),
            }

            from doc_helpers import safe_replace_in_doc
            for (file_path,) in documents:
                if not os.path.exists(file_path):
                    continue
                try:
                    doc = Document(file_path)
                    safe_replace_in_doc(doc, placeholders)
                    doc.save(file_path)
                    print(f"Successfully updated document: {file_path}")
                except Exception as doc_error:
                    print(f"Error updating document {file_path}: {doc_error}")

            QtWidgets.QApplication.restoreOverrideCursor()
            QMessageBox.information(self, "نجاح", "تم حفظ الجلسة وتحديث ملف الدعوى بنجاح ✅")
            self.show_scheduling()
        except Exception as e:
             if 'db' in locals():
                 try: db.close()
                 except: pass
             QtWidgets.QApplication.restoreOverrideCursor()
             QMessageBox.critical(self, "خطأ", f"حدث خطأ أثناء الحفظ: {e}")

    def new_case_dialog(self):
        db = DataBase()
        db.cur.execute("""
            SELECT client_id, plaintiff_name, defendant_name, case_type
            FROM cms.client
            WHERE client_id NOT IN (SELECT client_id FROM cms.case_client)
            """)
        clients = db.cur.fetchall()
        db.close()

        if not clients:
            QMessageBox.information(self, "تنبيه", "لا توجد عملاء في قاعدة البيانات بعد.")
            return

        dialog = QtWidgets.QDialog(self)
        dialog.setWindowTitle("إنشاء قضية جديدة")
        dialog.setFixedSize(550, 350)
        layout = QtWidgets.QVBoxLayout(dialog)
        dialog.setLayoutDirection(Qt.RightToLeft)

        layout.addWidget(QLabel("اختر القضايا:"))
        table = QtWidgets.QTableWidget()
        table.setRowCount(len(clients))
        table.setColumnCount(4)
        table.setHorizontalHeaderLabels(["اختيار", "المدعي", "المدعى عليه", "نوع القضية"])
        table.horizontalHeader().setSectionResizeMode(QtWidgets.QHeaderView.Stretch)
        table.setLayoutDirection(Qt.RightToLeft)

        checkboxes = []
        for row, client in enumerate(clients):
            client_id, p_name, d_name, c_type = client
            checkbox = QCheckBox()
            checkbox_widget = QWidget()
            cb_layout = QHBoxLayout(checkbox_widget)
            cb_layout.addWidget(checkbox)
            cb_layout.setAlignment(Qt.AlignCenter)
            cb_layout.setContentsMargins(0, 0, 0, 0)
            table.setCellWidget(row, 0, checkbox_widget)
            checkboxes.append(checkbox)
            table.setItem(row, 1, QtWidgets.QTableWidgetItem(p_name))
            table.setItem(row, 2, QtWidgets.QTableWidgetItem(d_name))
            table.setItem(row, 3, QtWidgets.QTableWidgetItem(c_type))

        layout.addWidget(table)
        save_btn = QPushButton("إنشاء القضايا")
        layout.addWidget(save_btn)

        def create_case():
            selected_rows = [i for i, cb in enumerate(checkboxes) if cb.isChecked()]
            if not selected_rows:
                QMessageBox.warning(dialog, "تنبيه", "اختر قضية واحدة على الأقل")
                return

            db = DataBase()
            for idx in selected_rows:
                client_id, p_name, d_name, c_type = clients[idx]
                db.cur.execute("SELECT count(*) FROM cms.court_case")
                case_count = db.cur.fetchone()[0]
                case_number = f"{datetime.now().strftime('%Y')}/{case_count + 1}"
                
                db.cur.execute("""
                    INSERT INTO cms.court_case (case_type, case_number, status, filing_date, year, description, created_by)
                    VALUES (%s, %s, %s, %s, %s, %s, %s)
                    RETURNING case_id
                """, (c_type, case_number, "مفتوحة", datetime.now().date(), datetime.now().year, "-", self.current_user_id))
                new_case_id = db.cur.fetchone()[0]

                db.cur.execute("INSERT INTO cms.case_client (case_id, client_id, role_in_case) VALUES (%s, %s, %s)", (new_case_id, client_id, "Plaintiff"))
                db.cur.execute("UPDATE cms.document SET case_id = %s WHERE client_id = %s AND case_id IS NULL", (new_case_id, client_id))

                # Notify Judges
                db.cur.execute("SELECT user_id FROM cms.users WHERE role_id = 4")
                judges_to_notify = db.cur.fetchall()
                for judge_entry in judges_to_notify:
                    notif_msg = f"تم إنشاء قضية جديدة: {c_type} برقم {case_number}"
                    db.cur.execute("INSERT INTO cms.notification (user_id, message) VALUES (%s, %s)", (judge_entry[0], notif_msg))

            db.conn.commit()
            db.close()
            QMessageBox.information(dialog, "نجاح", f"تم إنشاء {len(selected_rows)} قضية بنجاح ✅")
            dialog.accept()

        save_btn.clicked.connect(create_case)
        dialog.exec_()