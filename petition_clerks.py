import os
import sys
import shutil
import random
from datetime import date, datetime
from docx import Document
from db import DataBase
from PyQt5 import uic, QtWidgets
from PyQt5.QtWidgets import QMainWindow, QMessageBox
from PyQt5.QtCore import Qt, QTimer

def resource_path(relative_path):
    try:
        base_path = sys._MEIPASS
    except Exception:
        base_path = os.path.abspath(".")
    return os.path.join(base_path, relative_path)

class Petition_Clerks(QMainWindow):
    def __init__(self, current_user_id, main_shell=None):
        super().__init__()
        self.db = DataBase()
        self.c_u_i = current_user_id
        self.main_shell = main_shell
        self.current_case_type = None
        self.current_template = None

        uic.loadUi(resource_path("petition_clerks2.ui"), self)
        self.formFrame.setVisible(False)
        self.welcome_label.setText("اختر لائحة دعوى من القائمة")
        self.welcome_label.setStyleSheet("color: #452829; font-size: 32px; font-weight: bold;")
        self.welcome_label.setMinimumWidth(600)
        self.welcome_label.setVisible(True)
        self.setStyleSheet("""
        * {
            font-family: "Alyamama";
            color: #452829;
        }
        #sideBar {
            background-color: #452829;
            min-width: 260px;
        }
        #sideBar QPushButton {
            color: #f3e8df;
            background-color: transparent;
            border: none;
            text-align: right;
            padding-right: 25px;
            font-size: 19px;
            font-weight: bold;
            min-height: 55px;
            border-right: 5px solid transparent;
        }
        #sideBar QPushButton:hover {
            background-color: rgba(0, 0, 0, 0.3);
            color: white;
        }
        #sideBar QPushButton[active="true"] {
            background-color: rgba(0, 0, 0, 0.3) !important;
            color: white !important;
            border-right: 5px solid #b08d57;
            font-weight: bold;
        }
        }
        """)
        
        self.logoutBtn.setFixedSize(230, 40)
        self.logoutBtn.setStyleSheet("""
            QPushButton {
                color: #452829;
                background-color: white;
                border-radius: 10px;
                font-weight: bold;
                font-family: "Alyamama";
                font-size: 16px;
                border: none;
            }
            QPushButton:hover {
                background-color: #e63946;
                color: white;
            }
        """)
        
        self.sendFile.clicked.connect(self.process_full_workflow)
        self.logoutBtn.clicked.connect(self.log_out)
        self.sendFile.setFocusPolicy(Qt.NoFocus)
        self.logoutBtn.setFocusPolicy(Qt.NoFocus)
        
        self.case_config = {
            "case1": {"label": "نفقة زوجة", "template": "لائحة دعوى نفقة زوجة.docx"},
            "case2": {"label": "عفش بيت", "template": "لائحة دعوى عفش بيت.docx"}, 
            "case3": {"label": "مهر مؤجل", "template": "لائحة دعوى مهر مؤجل.docx"},
            "case4": {"label": "نفقة عفش غيابي", "template": "nafqa.docx"},
            "case5": {"label": "نفقة زوجة غيابي", "template": "nafqa.docx"},
            "case6": {"label": "نفقة صغار", "template": "nafqa.docx"},
        }

        for btn_id in self.case_config.keys():
            if hasattr(self, btn_id):
                btn = getattr(self, btn_id)
                btn.clicked.connect(self.handle_case_selection_click)
                btn.setFocusPolicy(Qt.NoFocus)

        self.load_receivers()

        if hasattr(self, 'employeesTable'):
             QTimer.singleShot(0, lambda: self.employeesTable.setFocus())

        # Tab Order Sequence
        self.setTabOrder(self.plaintiff_name, self.plaintiff_national_id)
        self.setTabOrder(self.plaintiff_national_id, self.plaintiff_phone)
        self.setTabOrder(self.plaintiff_phone, self.plaintiff_address)
        
        self.setTabOrder(self.plaintiff_address, self.defendant_name)
        self.setTabOrder(self.defendant_name, self.defendant_national_id)
        self.setTabOrder(self.defendant_national_id, self.defendant_phone)
        self.setTabOrder(self.defendant_phone, self.defendant_address)
        
        self.setTabOrder(self.defendant_address, self.divorce_date)
        self.setTabOrder(self.divorce_date, self.divorce_court)
        self.setTabOrder(self.divorce_court, self.iddah_end_date)
        self.setTabOrder(self.iddah_end_date, self.pregnancy_end_date)
        self.setTabOrder(self.pregnancy_end_date, self.furniture_value)
        self.setTabOrder(self.furniture_value, self.dowry_value)
        
        self.setTabOrder(self.dowry_value, self.comboBox)
        self.setTabOrder(self.comboBox, self.sendFile)

    def load_receivers(self):
        try:
            db = DataBase()
            db.cur.execute("SELECT user_id, full_name FROM cms.users WHERE role_id = '2'")
            users = db.cur.fetchall()
            self.comboBox.clear()
            self.comboBox.addItem("اختر الموظف المستلم", None)
            for user_id, name in users:
                self.comboBox.addItem(name, user_id)
            db.close()
        except Exception as e:
            QMessageBox.warning(self, "Error", f"Failed to load receivers: {str(e)}")

    def reset_sidebar_styles(self):
        for btn_id in self.case_config.keys():
            if hasattr(self, btn_id):
                btn = getattr(self, btn_id)
                btn.setProperty("active", False)
                btn.style().unpolish(btn)
                btn.style().polish(btn)

    def set_active_button(self, button):
        self.reset_sidebar_styles()
        button.setProperty("active", True)
        button.style().unpolish(button)
        button.style().polish(button)

    def handle_case_selection_click(self):
        btn = self.sender()
        if btn:
            self.set_active_button(btn)
            btn_id = btn.objectName()
            if btn_id in self.case_config:
                config = self.case_config[btn_id]
                self.current_case_type = config["label"]
                self.current_template = config["template"]
                self.label_2.setText(f"أدخل بيانات لائحة دعوى {self.current_case_type}")

                # Show form and hide welcome label
                self.welcome_label.setVisible(False)
                self.formFrame.setVisible(True)

                is_furniture = (self.current_case_type == "عفش بيت")
                is_dowry = (self.current_case_type == "مهر مؤجل")
                
                # Section Header Visibility
                self.hdr_divorce.setVisible(is_furniture or is_dowry)
                
                self.divorce_date.setVisible(is_furniture or is_dowry)
                self.divorce_court.setVisible(is_furniture or is_dowry)
                self.furniture_value.setVisible(is_furniture)
                
                self.iddah_end_date.setVisible(is_dowry)
                self.pregnancy_end_date.setVisible(is_dowry)
                self.dowry_value.setVisible(is_dowry)

    def process_full_workflow(self):
        if not hasattr(self, 'current_case_type') or not self.current_case_type:
            QMessageBox.warning(self, "تنبيه", "يرجى اختيار نوع القضية أولاً من القائمة الجانبية.")
            return

        receiver_index = self.comboBox.currentIndex()
        if receiver_index <= 0:
            QMessageBox.warning(self, "تنبيه", "يرجى اختيار الموظف المستلم.")
            return
        receiver_id = self.comboBox.itemData(receiver_index)

        try:
            db = DataBase()
            db.cur.execute("""
                INSERT INTO cms.client (plaintiff_name, plaintiff_national_id, plaintiff_phone, plaintiff_address,
                                        defendant_name, defendant_national_id, defendant_phone, defendant_address, case_type)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
                RETURNING client_id
            """, (
                self.plaintiff_name.text(), self.plaintiff_national_id.text(), self.plaintiff_phone.text(),
                self.plaintiff_address.text(), self.defendant_name.text(), self.defendant_national_id.text(),
                self.defendant_phone.text(), self.defendant_address.text(), self.current_case_type 
            ))
            client_id = db.cur.fetchone()[0]
            db.conn.commit()
            db.close()
        except Exception as e:
             QMessageBox.warning(self, "Error", f"Failed to register client: {str(e)}")
             return

        try:
            template_path = os.path.join("Template files", self.current_template)
            if not os.path.exists(template_path):
                QMessageBox.warning(self, "Error", f"Template file not found: {template_path}")
                return

            safe_name = "".join([c for c in self.plaintiff_name.text() if c.isalnum() or c in (' ', '_')]).rstrip()
            final_filename = f"{self.current_case_type} - {safe_name}.docx"
            final_file = os.path.join(self.db.files_path, final_filename)
            shutil.copy2(template_path, final_file)
            doc = Document(final_file)

            days_map = {
                "Monday": "الإثنين", "Tuesday": "الثلاثاء", "Wednesday": "الأربعاء",
                "Thursday": "الخميس", "Friday": "الجمعة", "Saturday": "السبت", "Sunday": "الأحد"
            }
            day_in_arabic = days_map.get(date.today().strftime("%A"), date.today().strftime("%A"))

            p_addr = self.plaintiff_address.text()
            p_from = p_addr.split('-')[0].strip() if '-' in p_addr else p_addr
            p_res = p_addr.split('-')[1].strip() if '-' in p_addr else "-"

            d_addr = self.defendant_address.text()
            d_from = d_addr.split('-')[0].strip() if '-' in d_addr else d_addr
            d_res = d_addr.split('-')[1].strip() if '-' in d_addr else "-"

            placeholders = {
                "{DATE_DAY}": day_in_arabic,
                "{DATE_FULL}": date.today().strftime("%d/%m/%Y"),
                "{PLAINTIFF_NAME}": self.plaintiff_name.text(),
                "{PLAINTIFF_FROM}": p_from,
                "{PLAINTIFF_RESIDENT}": p_res,
                "{DEFENDANT_NAME}": self.defendant_name.text(),
                "{DEFENDANT_FROM}": d_from,
                "{DEFENDANT_RESIDENT}": d_res,             
            }
            
            if self.current_case_type == "عفش بيت":
                placeholders["{DIVORCE_DATE}"] = self.divorce_date.text()
                placeholders["{FURNITURE_VALUE}"] = self.furniture_value.text()
                placeholders["{DIVORCE_COURT}"] = self.divorce_court.text()

            if self.current_case_type == "مهر مؤجل":
                placeholders["{DIVORCE_DATE}"] = self.divorce_date.text()
                placeholders["{DIVORCE_COURT}"] = self.divorce_court.text()
                placeholders["{DOWRY_VALUE}"] = self.dowry_value.text()
                
                # Extract first name from full name
                d_full_name = self.defendant_name.text().strip()
                d_first_name = d_full_name.split(' ')[0] if d_full_name else ""
                placeholders["{DEFENDANT_FIRST}"] = d_first_name
                
                # Logic for empty date fields: replace with "/  /" if empty
                iddah_val = self.iddah_end_date.text().strip()
                preg_val = self.pregnancy_end_date.text().strip()
                
                placeholders["{IDDAH_DATE}"] = iddah_val if iddah_val else "/  /"
                placeholders["{PREG_DATE}"] = preg_val if preg_val else "/  /"

            from doc_helpers import safe_replace_in_doc
            
            safe_replace_in_doc(doc, placeholders)
            doc.save(final_file)
            
            db = DataBase()
            db.cur.execute("""
                INSERT INTO cms.document (document_type, file_path, uploaded_by, client_id)
                VALUES (%s, %s, %s, %s)
                RETURNING document_id
            """, (self.current_case_type, final_file, self.c_u_i, client_id))
            doc_id = db.cur.fetchone()[0]
            db.conn.commit()
            db.close()
            
        except Exception as e:
            QMessageBox.warning(self, "Error", f"Failed to generate file: {str(e)}")
            return

        try:
            db = DataBase()
            db.cur.execute("""
                INSERT INTO cms.file_transfer (transfer_id, transfer_date, status, document_id, sender_id, receiver_id)
                VALUES (%s, %s, %s, %s, %s, %s)
            """, (random.randint(1, 1000000), date.today(), "pending", doc_id, self.c_u_i, receiver_id))
            db.conn.commit()
            db.close()
        except Exception as e:
            QMessageBox.warning(self, "Error", f"Failed to send file: {str(e)}")
            return

        try:
            notification_msg = f"({self.current_case_type} - \"{self.plaintiff_name.text().strip()}\" جديد)"
            db = DataBase()
            db.cur.execute("""
                INSERT INTO cms.notification (message, user_id, created_at, document_id)
                VALUES (%s, %s, NOW(), %s)
            """, (notification_msg, receiver_id, doc_id))
            db.conn.commit()
            db.close()
        except Exception as e:
            print(f"Failed to send notification: {e}")

        QMessageBox.information(self, "نجاح", f"تم تسجيل الموكل، إنشاء الملف وإرساله بنجاح!")
        
        self.plaintiff_name.clear()
        self.plaintiff_national_id.clear()
        self.plaintiff_phone.clear()
        self.plaintiff_address.clear()
        self.defendant_name.clear()
        self.defendant_national_id.clear()
        self.defendant_phone.clear()
        self.defendant_address.clear()
        self.divorce_date.clear()
        self.furniture_value.clear()
        self.divorce_court.clear()
        self.divorce_date.setVisible(False)
        self.furniture_value.setVisible(False)
        self.divorce_court.setVisible(False)
        self.formFrame.setVisible(False)
        self.welcome_label.setVisible(True)
        self.current_case_type = None
        self.reset_sidebar_styles()
        self.comboBox.setCurrentIndex(0)
        self.label_2.setText("أدخل بيانات لائحة دعوى:")

    def log_out (self):
        if self.main_shell:
            self.main_shell.switch_to_login()
        else:
            self.close()
